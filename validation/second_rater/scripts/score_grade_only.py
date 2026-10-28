# -*- coding: utf-8 -*-
"""10/21: 따로 받은 등급 20편 시트 채점. 우리 등급의 두 사실(code, measured; gradeval/answer_key_grade.csv)과 2차 평가자 Q1·Q2 를 사실별로,
   그리고 등급(A = 코드+측정, B = 측정만, C = 측정 없음)으로 비교. score_second_rater.py 의 등급 부분과 같은 규칙.
   출력: second_rater_grade_result.json. 실행: python main_survey/second_rater/score_grade_only.py <등급 시트>"""
import csv, json, os, sys
sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__)); MS = os.path.dirname(HERE)
from openpyxl import load_workbook
yn = lambda v: True if str(v or "").strip().startswith("예") else (False if str(v or "").strip().startswith("아니") else None)
def kappa(x, y):
    n = len(x); cats = sorted(set(x) | set(y), key=str); po = sum(a == b for a, b in zip(x, y)) / n
    pe = sum((x.count(c) / n) * (y.count(c) / n) for c in cats); return (round((po - pe) / (1 - pe), 2) if pe < 1 else 1.0), round(100 * po, 1)
KEY = [k for k in csv.DictReader(open(os.path.join(HERE, "key_second_rater.csv"), encoding="utf-8-sig")) if k["sheet"] == "G"]
KG = {int(k["row"]): k for k in csv.DictReader(open(os.path.join(MS, "gradeval", "answer_key_grade.csv"), encoding="utf-8-sig"))}
ws = load_workbook(sys.argv[1], data_only=True)["등급"]; head = [c.value for c in ws[1]]; col = head.index
S = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
gr = lambda c, m: "A" if c and m else ("B" if m else "C")
rows = []
for k in KEY:
    n, o = int(k["new_row"]), int(k["orig_row"]); g = KG[o]; s = S[n]
    c2, m2 = yn(s[col("Q1 코드 공개 (예/아니오)")]), yn(s[col("Q2 측정 비용 (예/아니오)")])
    c1, m1 = g["code"] == "True", g["measured"] == "True"
    rows.append(dict(n=n, name=g["name"], arxiv=g["arxiv_id"], ours=g["grade"], ours_code=c1, ours_meas=m1, r2_code=c2, r2_meas=m2, r2_grade=gr(c2, m2),
                     r2_code_note=str(s[col("Q1 메모(링크)")] or ""), r2_note=str(s[col("메모")] or "")))
res = {}
for tag, a, b in (("code", "ours_code", "r2_code"), ("measured", "ours_meas", "r2_meas"), ("grade", "ours", "r2_grade")):
    res[f"{tag}_kappa"], res[f"{tag}_agree"] = kappa([r[a] for r in rows], [r[b] for r in rows])
res["n"] = len(rows); res["disagreements"] = [r for r in rows if r["ours"] != r["r2_grade"] or r["ours_code"] != r["r2_code"] or r["ours_meas"] != r["r2_meas"]]
json.dump(res, open(os.path.join(HERE, "second_rater_grade_result.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print({k: v for k, v in res.items() if k != "disagreements"})
for r in res["disagreements"]:
    print(f"  #{r['n']} {r['name']} ({r['arxiv']}): 우리 {r['ours']} (code {r['ours_code']}, meas {r['ours_meas']}) / 2차 {r['r2_grade']} (code {r['r2_code']}, meas {r['r2_meas']})\n     링크메모: {r['r2_code_note'][:110]}\n     메모: {r['r2_note'][:140]}")
