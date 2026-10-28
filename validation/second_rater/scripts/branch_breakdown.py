# -*- coding: utf-8 -*-
"""2차 평가자 판정의 가지별 일치도와, 겹치는 표본에서 1차·2차 평가자 각각으로 본 엄격 규칙 정밀도(10/20).
   새 판정 없음 — 두 평가자의 기존 판정만 비교. 출력: second_rater_branch.json. 실행: python main_survey/second_rater/branch_breakdown.py <2차 평가자 시트>"""
import csv, json, os, sys
sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__)); MS = os.path.dirname(HERE)
from openpyxl import load_workbook
yn = lambda v: True if str(v or "").strip().startswith("예") else (False if str(v or "").strip().startswith("아니") else None)
def sheet(ws):
    head = [c.value for c in ws[1]]; return {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}, (lambda n: head.index(n))
KEY = [k for k in csv.DictReader(open(os.path.join(HERE, "key_second_rater.csv"), encoding="utf-8-sig")) if k["sheet"] == "A"]
KP = {int(k["row"]): k for k in csv.DictReader(open(os.path.join(MS, "classval", "answer_key_precision.csv"), encoding="utf-8-sig"))}
R1, c1 = sheet(load_workbook(os.path.join(MS, "classval", "returned_평가자1.xlsx"), data_only=True)["A_정밀도"])
R2, c2 = sheet(load_workbook(sys.argv[1], data_only=True)["A"])
out = {}
for k in KEY:
    n, o = int(k["new_row"]), int(k["orig_row"]); meta = KP[o]; b = meta["branch"]
    a1 = (yn(R1[o][c1("Q1 다루는가")]), yn(R1[o][c1("Q2 엄격")])); a2 = (yn(R2[n][c2("Q1 다루는가")]), yn(R2[n][c2("Q2 엄격")]))
    e = out.setdefault(b, dict(n=0, q1_agree=0, q2_agree=0, strict_n=0, strict_yes_r1=0, strict_yes_r2=0, strata={}))
    e["n"] += 1; e["q1_agree"] += a1[0] == a2[0]; e["q2_agree"] += a1[1] == a2[1]
    e["strata"][meta["stratum"]] = e["strata"].get(meta["stratum"], 0) + 1
    if meta["stratum"] == "strict":
        e["strict_n"] += 1; e["strict_yes_r1"] += bool(a1[1]); e["strict_yes_r2"] += bool(a2[1])
for b in sorted(out):
    e = out[b]; print(f"{b}: n={e['n']} strata={e['strata']} Q1 일치 {e['q1_agree']}/{e['n']}, Q2 일치 {e['q2_agree']}/{e['n']} | 엄격 층 {e['strict_n']}편: 엄격 '예' 1차 {e['strict_yes_r1']}, 2차 {e['strict_yes_r2']}")
json.dump(out, open(os.path.join(HERE, "second_rater_branch.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
