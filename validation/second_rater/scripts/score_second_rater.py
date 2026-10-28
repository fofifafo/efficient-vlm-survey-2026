# -*- coding: utf-8 -*-
"""2차 평가자 부분 이중 판정 채점. 시트 A·B 는 1차 평가자(classval/returned_평가자1.xlsx)의 같은 행과, 시트 C 는 1차 평가자의 새 표본 판정
   (두 번째 인자; 원문·한글본 어느 쪽이든), 등급은 우리 등급(gradeval/answer_key_grade.csv)과 비교.
   10/16 수정(종합 검토 지적): 열을 위치가 아니라 머리글 이름으로 찾음(한글본은 Q1·Q2 가 9·10열), 등급 시트가 없으면 건너뜀,
   빈 칸은 불일치로 세지 않고 제외.
   출력: second_rater_result.json, second_rater_disagreements.xlsx(불일치 목록).
   실행: python main_survey/second_rater/score_second_rater.py <채운 시트.xlsx> [1차 평가자 새 표본 시트.xlsx]"""
import csv, json, os, re, sys
sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__)); MS = os.path.dirname(HERE)
from openpyxl import load_workbook, Workbook
yn = lambda v: True if str(v or "").strip().startswith("예") else (False if str(v or "").strip().startswith("아니") else None)
codes = lambda v: None if str(v or "").strip() == "" else (set() if "없" in str(v) else set(re.findall(r"4\.[1-6]", str(v))))
BR = ("4.1", "4.2", "4.3", "4.4", "4.5", "4.6")
def kappa(x, y):
    n = len(x)
    if not n: return None, None
    cats = sorted(set(x) | set(y), key=str); po = sum(a == b for a, b in zip(x, y)) / n
    pe = sum((x.count(c) / n) * (y.count(c) / n) for c in cats); return (round((po - pe) / (1 - pe), 2) if pe < 1 else 1.0), round(100 * po, 1)
def sheet(ws):
    head = [c.value for c in ws[1]]
    col = lambda *names: next(head.index(n) for n in names if n in head)
    return {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}, col
KEY = list(csv.DictReader(open(os.path.join(HERE, "key_second_rater.csv"), encoding="utf-8-sig")))
r1 = load_workbook(os.path.join(MS, "classval", "returned_평가자1.xlsx"), data_only=True)
R1A, c1A = sheet(r1["A_정밀도"]); R1B, c1B = sheet(r1["B_재현율"])
KG = {int(k["row"]): k for k in csv.DictReader(open(os.path.join(MS, "gradeval", "answer_key_grade.csv"), encoding="utf-8-sig"))}
s2 = load_workbook(sys.argv[1], data_only=True)
res, dis = {"sheets": s2.sheetnames}, []
def compare_yn(tag, m, R1, c1, S2, c2):
    for q, names in (("Q1", ("Q1 다루는가",)), ("Q2", ("Q2 엄격",))):
        x = [yn(R1[o][c1(*names)]) for n, o in m]; y = [yn(S2[n][c2(*names)]) for n, o in m]
        pairs = [((n, o), a, b) for (n, o), a, b in zip(m, x, y) if a is not None and b is not None]
        res[f"{tag}_{q}_kappa"], res[f"{tag}_{q}_agree"] = kappa([a for _, a, _ in pairs], [b for _, _, b in pairs]); res[f"{tag}_{q}_n"] = len(pairs)
        dis.extend((tag, n, R1[o][1], R1[o][2], q, a, b) for (n, o), a, b in pairs if a != b)
S2A, c2A = sheet(s2["A"])
compare_yn("A", [(int(k["new_row"]), int(k["orig_row"])) for k in KEY if k["sheet"] == "A"], R1A, c1A, S2A, c2A)
mC = [(int(k["new_row"]), int(k["orig_row"])) for k in KEY if k["sheet"] == "C"]
if mC and len(sys.argv) > 2 and "C" in s2.sheetnames:
    R1C, c1C = sheet(load_workbook(sys.argv[2], data_only=True)["A_정밀도"]); S2C, c2C = sheet(s2["C"])
    compare_yn("C", mC, R1C, c1C, S2C, c2C)
S2B, c2B = sheet(s2["B"])
mB = [(int(k["new_row"]), int(k["orig_row"])) for k in KEY if k["sheet"] == "B"]
for lv, names in (("loose", ("다루는 분야 (예: 4.2, 4.4 / 없음)",)), ("strict", ("엄격까지 만족하는 분야",))):
    x, y = [], []
    for n, o in mB:
        a, b = codes(R1B[o][c1B(*names)]), codes(S2B[n][c2B(*names)])
        if a is None or b is None: continue
        x += [br in a for br in BR]; y += [br in b for br in BR]
    res[f"B_{lv}_kappa"], res[f"B_{lv}_agree"] = kappa(x, y); res[f"B_{lv}_n_papers"] = len(x) // 6
if "등급" in s2.sheetnames:
    S2G, c2G = sheet(s2["등급"])
    gr = lambda c, m: "A" if c and m else ("B" if m else "C")
    pairs = []
    for n, o in [(int(k["new_row"]), int(k["orig_row"])) for k in KEY if k["sheet"] == "G"]:
        q1, q2 = yn(S2G[n][c2G("Q1 코드 공개 (예/아니오)")]), yn(S2G[n][c2G("Q2 측정 비용 (예/아니오)")])
        if q1 is None or q2 is None: continue
        pairs.append((n, o, KG[o]["grade"], gr(q1, q2)))
    res["G_kappa"], res["G_agree"] = kappa([p[2] for p in pairs], [p[3] for p in pairs]); res["G_n"] = len(pairs)
    dis.extend(("G", n, KG[o]["name"], KG[o]["arxiv_id"], "grade", a, b) for n, o, a, b in pairs if a != b)
else:
    res["G_skipped"] = "등급 시트 없음"
wb = Workbook(); ws = wb.active; ws.append(["시트", "번호", "분야/이름", "arXiv", "질문", "1차(또는 우리 등급)", "2차", "합의"])
for d in dis: ws.append(list(d) + [""])
wb.save(os.path.join(HERE, "second_rater_disagreements.xlsx")); res["n_disagreements"] = len(dis)
json.dump(res, open(os.path.join(HERE, "second_rater_result.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(json.dumps(res, ensure_ascii=False, indent=1))
