# -*- coding: utf-8 -*-
"""10/20: 2차 평가자에게서 돌아온 시트에 '등급' 시트가 없어(보낸 파일에 빠져 있었던 것으로 보임) 등급 20편만 따로 보낸다.
   2차평가자_판정시트_v2.xlsx 에서 '등급' 시트만 남기고 안내를 등급용으로 바꾼다. 정답(우리 등급)·원래 행 번호는 들어 있지 않다
   (원래 행 번호는 key_second_rater.csv 의 G 행에만). 채점은 score_second_rater.py 가 '등급' 시트를 그대로 읽는다.
   출력: to_rater/20261020_등급_판정_시트.xlsx. 실행: python main_survey/second_rater/make_grade_only_sheet.py"""
import os, sys
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
sys.stdout.reconfigure(encoding="utf-8"); HERE = os.path.dirname(os.path.abspath(__file__))
wb = load_workbook(os.path.join(HERE, "2차평가자_판정시트_v2.xlsx"))
for name in ("기준", "A", "B", "C"): del wb[name]
an = wb["안내"]; an.delete_rows(1, an.max_row)
TXT = ["효율적 VLM 서베이 — 2차 평가자 등급 판정 시트 (20편). 도와주셔서 감사합니다. 약 1시간 걸립니다.",
       "",
       "지난번 시트(A·B·C)에 이 '등급' 시트가 빠져 있어서 따로 보내 드립니다.",
       "목적: 논문마다 매긴 근거 등급이 맞는지 사람이 독립적으로 확인합니다. 다른 사람(제1저자 포함)의 판정은 보지 않고, 상의하지 않고, AI 도구 없이 직접 판단해 주세요.",
       "",
       "['등급' 시트 — 20편] arXiv 링크로 논문을 열어 두 가지만 확인합니다.",
       "   Q1 코드 공개: 논문에 코드 저장소 링크가 있고, 열었을 때 실제 코드가 있으면 예. 링크가 없거나 README만 있으면(\"coming soon\") 아니오.",
       "      모델 가중치와 실행 코드를 Hugging Face 등에 공개했으면 예. 확인한 링크를 'Q1 메모(링크)' 칸에 적어 주세요.",
       "   Q2 측정 비용: 저자가 자기 시스템에서 직접 잰 시간(지연·속도 향상 배수·처리량·학습 시간) 또는 메모리 수치가 본문·표·그림에 있으면 예.",
       "      FLOPs·파라미터 수·토큰 수만 있거나 다른 논문 수치를 인용한 것이면 아니오. 예이면 근거 위치(예: 표 3)를 'Q2 근거 위치' 칸에 적어 주세요.",
       "",
       "애매하면 메모 칸에 이유를 짧게 적어 주세요. 다 하시면 파일을 저장해서 보내 주세요."]
for i, t in enumerate(TXT, 1):
    c = an.cell(row=i, column=1, value=t); c.alignment = Alignment(wrap_text=False)
    if i == 1: c.font = Font(bold=True, size=12)
an.column_dimensions["A"].width = 150
gs = wb["등급"]; head = [c.value for c in gs[1]]
assert head == ["번호", "논문", "arXiv 링크", "Q1 코드 공개 (예/아니오)", "Q1 메모(링크)", "Q2 측정 비용 (예/아니오)", "Q2 근거 위치", "메모"], head
assert gs.max_row == 21 and all(gs.cell(r, c).value in (None, "") for r in range(2, 22) for c in (4, 5, 6, 7, 8)), "답 칸이 비어 있어야 함"
wb.active = 0
out = os.path.join(HERE, "to_rater", "20261020_등급_판정_시트.xlsx"); wb.save(out)
print("saved", out, wb.sheetnames)
