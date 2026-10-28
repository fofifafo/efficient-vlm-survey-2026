# -*- coding: utf-8 -*-
"""채운 추출 대조 시트 집계 → extraction_check_result.json. 영문판·한글판 시트 모두 머리글 이름으로 칸을 찾는다.
   사용: python main_survey/extraction_check/score_extraction_check.py <시트.xlsx>"""
import collections, json, os, sys
from openpyxl import load_workbook
sys.stdout.reconfigure(encoding="utf-8"); HERE = os.path.dirname(os.path.abspath(__file__))
ws = load_workbook(sys.argv[1])["대조"]
H = [str(c.value or "") for c in ws[1]]
col = lambda name: next(i for i, h in enumerate(H) if h.startswith(name))
iv, ifl, im, ival, inote = col("판정"), col("칸"), col("방법"), col("추출값"), col("메모")
rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
c = collections.Counter(str(r[iv] or "").strip() for r in rows)
out = dict(n=len(rows), matched=c["일치"], partial=c["부분 일치"], mismatched=c["불일치"], blank=c[""],
           by_field={f: dict(collections.Counter(str(r[iv] or "") for r in rows if r[ifl] == f)) for f in {r[ifl] for r in rows}},
           to_fix=[dict(no=r[0], method=r[im], field=r[ifl], value=r[ival], verdict=r[iv], note=r[inote]) for r in rows if str(r[iv] or "") in ("부분 일치", "불일치")])
json.dump(out, open(os.path.join(HERE, "extraction_check_result.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(json.dumps(out, ensure_ascii=False, indent=1))
