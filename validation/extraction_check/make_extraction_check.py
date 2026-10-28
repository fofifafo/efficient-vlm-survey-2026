# -*- coding: utf-8 -*-
"""추출 수치 30개 원문 대조 시트 생성 (2026-10-12, 지유찬 결정: 본인이 대조).
   표본: 표 3–7에 들어간 영문 칸(matrix_en.csv) 중 숫자가 있는 칸. 등급에 쓰이는 속도·지연·메모리 칸 20개 + 감소량·FLOPs 칸 10개, 고정 시드.
   출력: 추출대조_시트.xlsx (판정 칸 비움), sample.csv (표본 기록)
   채점: python main_survey/extraction_check/score_extraction_check.py <채운 시트.xlsx>
   실행: python main_survey/extraction_check/make_extraction_check.py"""
import csv, glob, json, os, random, re, sys
sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.abspath(os.path.join(HERE, "..", "..")); os.chdir(ROOT); sys.path.insert(0, ROOT)
SEED = 20261012
import arxiv_scrape as A
corpus, _ = A.load_cache()
name = {}
for f in glob.glob("main_survey/matrix_4.*.csv"):
    for r in csv.DictReader(open(f, encoding="utf-8-sig")): name[(os.path.basename(f)[7:10], r["id"])] = r["name"]
ev = {}
for f in glob.glob("main_survey/matrix_evidence_*.json"):
    for x in json.load(open(f, encoding="utf-8")): ev.setdefault(x["id"], {}).update(x.get("evidence", {}))
EN = list(csv.DictReader(open("main_survey/matrix_en.csv", encoding="utf-8-sig")))
DASH = lambda v: not v.strip() or v.strip().startswith(("—", "-", "–")) and not re.search(r"\d", v)
FIELD = {"speed_latency": "속도·지연", "memory": "메모리", "reduction": "감소량·설정", "flops": "FLOPs"}
HINT = {"speed_latency": ("SPEED", "LAT", "THR"), "memory": ("MEM",), "reduction": ("TOK", "BIT", "PARAM"), "flops": ("FLOP",)}
cells = lambda fields: [(r, f) for r in EN for f in fields if not DASH(r[f]) and re.search(r"\d", r[f])]
rng = random.Random(SEED)
sample = rng.sample(cells(("speed_latency", "memory")), 20) + rng.sample(cells(("reduction", "flops")), 10)
rng.shuffle(sample)
with open(os.path.join(HERE, "sample.csv"), "w", encoding="utf-8-sig", newline="") as fo:
    w = csv.writer(fo); w.writerow(["no", "branch", "id", "field", "value"])
    for i, (r, f) in enumerate(sample, 1): w.writerow([i, r["branch"], r["id"], f, r[f]])

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
wb = Workbook(); ws = wb.active; ws.title = "대조"
ws.append(["번호", "가지", "방법", "논문 제목", "arXiv", "칸", "추출값 (표에 들어간 값)", "참고 문장 (자동 수집, 없거나 틀릴 수 있음)", "판정", "메모"])
for i, (r, f) in enumerate(sample, 1):
    hints = [s for k in HINT[f] for s in ev.get(r["id"], {}).get(k, [])][:2]
    title = " ".join(corpus.get(r["id"], {}).get("title", "").split())
    ws.append([i, r["branch"], name.get((r["branch"], r["id"]), ""), title, f"https://arxiv.org/abs/{r['id']}", FIELD[f], r[f],
               "\n\n".join(" ".join(h.split())[:400] for h in hints), "", ""])
    ws.cell(i + 1, 5).hyperlink = f"https://arxiv.org/abs/{r['id']}"; ws.cell(i + 1, 5).font = Font(color="0563C1", underline="single")
thin = Side(style="thin", color="BFBFBF")
for j, wdt in enumerate([6, 6, 18, 36, 14, 11, 36, 70, 12, 24], 1): ws.column_dimensions[chr(64 + j)].width = wdt
for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E7E6E6")
for rr in range(2, ws.max_row + 1): ws.cell(rr, 9).fill = PatternFill("solid", fgColor="FFF2CC"); ws.row_dimensions[rr].height = 120
dv = DataValidation(type="list", formula1='"일치,부분 일치,불일치"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"I2:I{ws.max_row}")
ws.freeze_panes = "D2"
g = wb.create_sheet("안내")
for line in ["추출 수치 30개 원문 대조 (논문 3.5절: 'X of 30 matched')", "",
             "1. arXiv 링크로 논문을 열고, G열 추출값의 숫자가 논문에 그대로 있는지 찾아 주세요. H열 참고 문장은 찾는 데 도움용이고, 판정은 반드시 원문으로 하세요.",
             "2. 판정 기준",
             "   일치: G열의 숫자가 모두 논문에 있고, 단위·설정(모델, 토큰 수, GPU 등)도 맞다.",
             "   부분 일치: 숫자는 맞지만 단위·설정 설명이 빠졌거나 일부 숫자를 원문에서 못 찾았다.",
             "   불일치: 숫자가 틀렸거나, 원문에 없거나, 뜻이 다르다(예: 감소율을 배율로 적음).",
             "3. 부분 일치나 불일치면 J열에 원문 문장이나 쪽 번호를 짧게 남겨 주세요. 고칠 때 씁니다.",
             "4. 다 하면 파일 경로를 Claude에게 알려 주세요. score_extraction_check.py 로 집계합니다.",
             f"표본: 속도·지연·메모리 칸 20개 + 감소량·FLOPs 칸 10개, 무작위(시드 {SEED})."]:
    g.append([line])
g.column_dimensions["A"].width = 130
wb.move_sheet("안내", offset=-1); wb.active = 0
wb.save(os.path.join(HERE, "추출대조_시트.xlsx"))
print("saved", len(sample), "cells; papers:", len({r['id'] for r, _ in sample}), "; with hints:", sum(1 for r, f in sample if any(ev.get(r['id'], {}).get(k) for k in HINT[f])))
