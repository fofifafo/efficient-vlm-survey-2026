# -*- coding: utf-8 -*-
"""Export the first rater's classification-validation judgments from the filled .xlsx sheets to CSV for release.
Titles and abstracts (and their machine translations) are dropped: abstracts are not redistributed; join on arXiv id.
Yes/no answers ("예"/"아니오") are mapped to 1/0. Rater notes are kept verbatim (mostly Korean).

Usage: python export_rater1_csv.py <returned_round1.xlsx> <returned_round2.xlsx> <out_dir>
Round 1 = rules v2 (sheets A precision 300 rows, B recall 100 rows); round 2 = rules v2.1 re-check (55 new rows, 4.3/4.5).
"""
import csv, os, sys
from openpyxl import load_workbook

def yn(v):
    s = str(v or "").strip()
    return 1 if s.startswith("예") else (0 if s.startswith("아니") else "")

def sheet(ws):
    rows = list(ws.iter_rows(values_only=True)); return [str(h) if h is not None else "" for h in rows[0]], [r for r in rows[1:] if r[0]]

def export_precision(ws, path):
    h, R = sheet(ws); c = {k: h.index(k) for k in ("번호", "분야", "arXiv", "Q1 다루는가", "Q2 엄격", "메모")}
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(["row", "branch", "arxiv_id", "q1_addresses_branch", "q2_strict", "rater_note"])
        for r in R:
            w.writerow([r[c["번호"]], str(r[c["분야"]])[:3], r[c["arXiv"]], yn(r[c["Q1 다루는가"]]), yn(r[c["Q2 엄격"]]), r[c["메모"]] or ""])
    return len(R)

def export_recall(ws, path):
    h, R = sheet(ws)
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(["row", "arxiv_id", "branches_addressed", "branches_strict", "rater_note"])
        for r in R:
            w.writerow([r[0], r[1], r[4] or "", r[5] or "", r[6] or ""])
    return len(R)

r1, r2, out = sys.argv[1:4]; os.makedirs(out, exist_ok=True)
wb = load_workbook(r1, data_only=True)
print("round1 precision", export_precision(wb["A_정밀도"], os.path.join(out, "round1_precision_rater1.csv")))
print("round1 recall", export_recall(wb["B_재현율"], os.path.join(out, "round1_recall_rater1.csv")))
wb = load_workbook(r2, data_only=True)
print("round2 precision", export_precision(wb["A_정밀도"], os.path.join(out, "round2_precision_rater1.csv")))
