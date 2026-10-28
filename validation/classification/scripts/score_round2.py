# -*- coding: utf-8 -*-
"""분류 규칙 v2.1(4.3·4.5) 재검증 채점(2026-10-15). 입력: 1차 평가자가 채운 새 표본 시트(원문 또는 한글본; 열은 이름으로 찾음).
   정밀도: 엄격 = 엄격 층의 Q2 비율, 느슨(4.3) = 엄격 층·느슨만 층의 Q1 비율을 층 크기(n_stratum)로 가중. 연도별·전체, Wilson 95% 구간.
   4.5 는 엄격 규칙만 바뀌었으므로 엄격만. 같은 표본으로 고친 1차 표본의 낙관적 추정과 섞지 않는다.
   실행: python main_survey/classval/score_round2.py <채운 시트.xlsx> → round2_result.json"""
import csv, json, math, os, sys
sys.stdout.reconfigure(encoding="utf-8"); HERE = os.path.dirname(os.path.abspath(__file__))
from openpyxl import load_workbook
K = {int(k["row"]): k for k in csv.DictReader(open(os.path.join(HERE, "answer_key_round2.csv"), encoding="utf-8-sig"))}
ws = load_workbook(sys.argv[1], data_only=True)["A_정밀도"]
head = [c.value for c in ws[1]]; c1, c2 = head.index("Q1 다루는가"), head.index("Q2 엄격")
yn = lambda v: True if str(v or "").strip().startswith("예") else (False if str(v or "").strip().startswith("아니") else None)
A = {int(r[0]): (yn(r[c1]), yn(r[c2])) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
def wilson(k, n):
    if not n: return (None, None)
    p = k / n; z = 1.96; d = 1 + z * z / n; c = (p + z * z / (2 * n)) / d; h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return round(100 * (c - h), 1), round(100 * (c + h), 1)
res = {"unfilled": [i for i, v in A.items() if None in v]}
for b in ("4.3", "4.5"):
    for scope in ("2024", "2026", "all"):
        rows = [(i, k) for i, k in K.items() if k["branch"] == b and (scope == "all" or k["year"] == scope) and None not in A[i]]
        s = [A[i][1] for i, k in rows if k["stratum"] == "strict"]
        out = dict(strict_n=len(s), strict_yes=sum(s), strict_prec=round(100 * sum(s) / len(s), 1) if s else None, strict_ci=wilson(sum(s), len(s)))
        if b == "4.3":
            num = den = 0.0; nl = 0
            for y in (("2024", "2026") if scope == "all" else (scope,)):
                for st in ("strict", "loose_only"):
                    R = [(i, k) for i, k in rows if k["year"] == y and k["stratum"] == st and k["loose_hit"] == "True"]
                    if not R: continue
                    v = [A[i][0] for i, _ in R]; Np = int(R[0][1]["n_stratum"]); num += Np * sum(v) / len(v); den += Np; nl += len(v)
            out.update(loose_prec=round(100 * num / den, 1) if den else None, loose_n=nl)
        res[f"{b}.{scope}"] = out
json.dump(res, open(os.path.join(HERE, "round2_result.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(json.dumps(res, ensure_ascii=False, indent=1))
