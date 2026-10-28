# -*- coding: utf-8 -*-
"""2차 리뷰 M1(10/14): 가지 분류 규칙(main_taxonomy)의 정밀도·재현율 검증 시트 — 두 평가자가 독립적으로 판정.
   A. 정밀도 표본(VLM 전체 코퍼스): 가지 6 × 연도 2(2024, 2026) × 25편 = 300편.
      층마다 엄격 일치에서 13편 + 느슨하게만 일치한 논문에서 12편(무작위, 시드 20261014). 평가자에게는 어느 규칙에 걸렸는지 보이지 않는다.
      질문 Q1(느슨) = 이 분야를 의도한 뜻으로 다루는가, Q2(엄격) = 이 분야의 방법을 제안하는가(4.5 실제 기기 실행·측정, 4.6 단위 붙은 수치).
   B. 재현율 표본(효율 풀): 2024 50편 + 2026 50편 = 100편, 규칙 결과와 무관하게 무작위. 평가자가 해당 가지 번호를 적는다.
   정답(규칙 결과·층)은 answer_key_*.csv 에만 있고 시트에는 없다. 채점: python main_survey/classval/score_classval.py <평가자1.xlsx> <평가자2.xlsx>
   출력: main_survey/classval/분류검증_시트_평가자1.xlsx, _평가자2.xlsx(내용 같음), answer_key_precision.csv, answer_key_recall.csv"""
import csv, os, random, sys
sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.abspath(os.path.join(HERE, "..", "..")); sys.path.insert(0, ROOT); os.chdir(ROOT)
_o = sys.stdout; sys.stdout = open(os.devnull, "w", encoding="utf-8")
import main_taxonomy as mt
sys.stdout = _o
SEED = 20261014; PER_STRICT, PER_LOOSE = 13, 12
NAME = {"4.1": "설계(구조)", "4.2": "학습(PEFT·데이터)", "4.3": "압축(양자화·저랭크·증류·프루닝)", "4.4": "실행(토큰·KV 캐시·프리필·디코딩)", "4.5": "배포(실제 기기)", "4.6": "측정(효율 지표)"}
Q1 = {"4.1": "MoE·조기 종료·층 축소·경량/무(無) 인코더·상태공간(Mamba)·선형 어텐션·Q-Former/리샘플러 같은 구조 효율 기법을 그 뜻으로 다룬다(제안·사용·비교·분석 모두 포함)",
      "4.2": "LoRA·PEFT·어댑터·프롬프트 튜닝·데이터 선별/효율 같은 학습 효율 기법을 그 뜻으로 다룬다",
      "4.3": "양자화·저랭크 분해·지식 증류·가중치/구조 프루닝·희소화 같은 모델 압축을 그 뜻으로 다룬다 (벡터 양자화 토크나이저, 데이터 증류, 행동 양자화는 아니오)",
      "4.4": "시각 토큰 가지치기/병합/압축·KV 캐시·프리필·추측 디코딩·서빙(vLLM 등) 같은 실행 효율 기법을 그 뜻으로 다룬다",
      "4.5": "실제 기기(Jetson·Snapdragon·스마트폰·NPU·FPGA·저전력 하드웨어 등)에서의 실행이나 에너지·전력 효율을 그 뜻으로 다룬다",
      "4.6": "지연시간·처리량·FLOPs·메모리·에너지 같은 효율 지표를 모델 비용의 뜻으로 다룬다 (예: 'latency'가 의학 용어로 쓰이면 아니오)"}
Q2 = {"4.1": "이 논문이 그런 구조 효율 기법(새 구조·변형)을 직접 제안한다 (기존 MoE 모델을 그냥 쓴 것은 아니오)",
      "4.2": "이 논문이 학습 효율 기법 자체를 제안한다 (LoRA를 도구로 써서 다른 과제를 푼 것은 아니오)",
      "4.3": "이 논문이 VLM 압축 방법을 제안한다",
      "4.4": "이 논문이 실행 효율 기법을 제안한다 (토큰 수를 언급만 한 것은 아니오)",
      "4.5": "초록에 이름 있는 실제 기기에서 실행했거나 측정했다고 적혀 있다 (\"엣지 배포 가능\" 같은 희망 서술만이면 아니오)",
      "4.6": "초록에 효율 지표의 단위 붙은 수치(ms, GB, FLOPs, ×배 빨라짐 등)가 있다"}
rng = random.Random(SEED)
yr = mt.yr
prec, key_p = [], []
for b, lab, loose, strict, _ in mt.RX:
    if b in mt.STRICT_OVERRIDE: S = {p["id"] for p in mt.core if strict.search(p["t"])}
    else: S = {p["id"] for p in mt.core if strict.search(p["t"]) or loose.search(p["title"])}
    L = {p["id"] for p in mt.core if loose.search(p["t"])}
    for y in ("2024", "2026"):
        sp = sorted(p["id"] for p in mt.core if yr(p) == y and p["id"] in S)
        lp = sorted(p["id"] for p in mt.core if yr(p) == y and p["id"] in L and p["id"] not in S)
        for st, pool, k in (("strict", sp, PER_STRICT), ("loose_only", lp, PER_LOOSE)):
            pick = rng.sample(pool, min(k, len(pool)))
            for x in pick:
                prec.append((b, x)); key_p.append(dict(branch=b, arxiv_id=x, year=y, stratum=st, n_stratum=len(pool), strict_hit=st == "strict", loose_hit=x in L))
order = list(range(len(prec))); rng.shuffle(order); order.sort(key=lambda i: prec[i][0])      # 가지끼리 묶되 가지 안 순서·연도는 섞음
prec = [prec[i] for i in order]; key_p = [key_p[i] for i in order]
for i, k in enumerate(key_p, 1): k["row"] = i
byid = {p["id"]: p for p in mt.papers}
poolY = {y: sorted(p["id"] for p in mt.pool if yr(p) == y) for y in ("2024", "2026")}
rec = rng.sample(poolY["2024"], 50) + rng.sample(poolY["2026"], 50); rng.shuffle(rec)
key_r = []
for i, x in enumerate(rec, 1):
    p = byid[x]; d = dict(row=i, arxiv_id=x, year=yr(p))
    for b, lab, loose, strict, _ in mt.RX:
        sh = bool(strict.search(p["t"])) if b in mt.STRICT_OVERRIDE else bool(strict.search(p["t"]) or loose.search(p["title"]))
        d[f"loose_{b}"] = bool(loose.search(p["t"])); d[f"strict_{b}"] = sh
    key_r.append(d)
for name, rows in (("answer_key_precision.csv", key_p), ("answer_key_recall.csv", key_r)):
    with open(os.path.join(HERE, name), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
thin = Side(style="thin", color="BFBFBF"); YEL = PatternFill("solid", fgColor="FFF2CC"); GRY = PatternFill("solid", fgColor="E7E6E6")
def style(ws, widths, fill_cols, h):
    for j, w in enumerate(widths, 1): ws.column_dimensions[chr(64 + j)].width = w
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for c in ws[1]: c.font = Font(bold=True); c.fill = GRY
    for r in range(2, ws.max_row + 1):
        ws.row_dimensions[r].height = h
        for c in fill_cols: ws.cell(r, c).fill = YEL
def build(path):
    wb = Workbook(); g = wb.active; g.title = "안내"
    for line in ["분류 규칙 검증 — 규칙이 논문을 가지(분야)에 맞게 붙였는지 사람이 확인합니다. 두 평가자가 서로 보지 않고 따로 판정합니다.", "",
                 "시트 A (정밀도, 300편): 각 행에 적힌 분야 하나에 대해서만 판단합니다. 제목과 초록만 봅니다.",
                 "  Q1 — 이 논문이 그 분야를 의도한 뜻으로 다루는가? (D열 기준)  →  예/아니오",
                 "  Q2 — 더 엄격한 기준(E열)을 만족하는가?  →  예/아니오.  Q1이 아니오면 Q2도 보통 아니오입니다.",
                 "  같은 논문이 다른 분야 행으로 또 나올 수 있습니다. 각 행은 그 행의 분야로만 판단하세요.",
                 "",
                 "시트 B (재현율, 100편): 분야가 정해져 있지 않습니다. 초록을 읽고",
                 "  E열 — 이 논문이 다루는 분야 번호를 모두 (예: 4.2, 4.4). 없으면 '없음'.  기준은 시트 A의 Q1과 같습니다.",
                 "  F열 — 그중 엄격 기준(시트 A의 Q2)까지 만족하는 분야 번호를 모두. 없으면 '없음'.",
                 "  분야별 기준은 '기준' 시트에 있습니다.",
                 "",
                 "규칙: 애매하면 기준에 더 가까운 쪽으로 고르고 메모 칸에 이유를 짧게. 다른 평가자와 상의하지 않습니다. answer_key 파일은 열지 않습니다.",
                 "다 하면 파일을 저장하고 Claude에게 경로를 알려 주세요. 예상 시간: 시트 A 약 3시간, 시트 B 약 1.5시간."]:
        g.append([line])
    g.column_dimensions["A"].width = 130
    k = wb.create_sheet("기준"); k.append(["분야", "Q1 (다루는가)", "Q2 (엄격)"])
    for b in NAME: k.append([f"{b} {NAME[b]}", Q1[b], Q2[b]])
    style(k, [26, 80, 80], [], 60)
    a = wb.create_sheet("A_정밀도"); a.append(["번호", "분야", "arXiv", "Q1 기준", "Q2 기준", "제목", "초록", "Q1 다루는가", "Q2 엄격", "메모"])
    for i, (b, x) in enumerate(prec, 1):
        p = byid[x]; a.append([i, f"{b} {NAME[b]}", x, Q1[b], Q2[b], " ".join(p["title"].split()), " ".join(p["abstract"].split()), "", "", ""])
    style(a, [6, 14, 12, 30, 30, 30, 90, 10, 10, 20], [8, 9], 200)
    dv = DataValidation(type="list", formula1='"예,아니오"', allow_blank=True); a.add_data_validation(dv); dv.add(f"H2:I{a.max_row}"); a.freeze_panes = "D2"
    r_ = wb.create_sheet("B_재현율"); r_.append(["번호", "arXiv", "제목", "초록", "다루는 분야 (예: 4.2, 4.4 / 없음)", "엄격까지 만족하는 분야", "메모"])
    for i, x in enumerate(rec, 1):
        p = byid[x]; r_.append([i, x, " ".join(p["title"].split()), " ".join(p["abstract"].split()), "", "", ""])
    style(r_, [6, 12, 30, 100, 18, 18, 20], [5, 6], 200); r_.freeze_panes = "C2"
    wb.active = 0; wb.save(path)
for who in ("평가자1", "평가자2"): build(os.path.join(HERE, f"분류검증_시트_{who}.xlsx"))
print("precision rows", len(prec), "| recall rows", len(rec))
import collections
print("strata (branch, year, stratum, picked/available):")
c = collections.Counter((k["branch"], k["year"], k["stratum"]) for k in key_p); av = {(k["branch"], k["year"], k["stratum"]): k["n_stratum"] for k in key_p}
for kk in sorted(c): print(" ", kk, f"{c[kk]}/{av[kk]}")
