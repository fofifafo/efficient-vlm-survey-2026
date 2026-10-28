# -*- coding: utf-8 -*-
"""4.3·4.5 규칙 수정안 시험(2026-10-14, 안 1). 기준 문서(SCREENING_CRITERIA.md) 제외 항목을 규칙에 구현:
   4.3 = 양자화는 모델 양자화만(벡터 양자화·양자화 토큰/토크나이저·코드북·행동 양자화 제외), 증류는 작은/압축 학생 문맥만, 희소화는 가중치·모델 희소화만.
   4.5 엄격 = 'Pixel' 은 구글 폰(Google Pixel, Pixel 숫자)만, 기기 이름 끝에 단어 경계.
   출력: 1차 평가자 판정으로 본 기존/수정 규칙 정밀도(같은 표본으로 고쳤으므로 낙관적 — 새 표본으로 다시 검증), 코퍼스 건수 변화."""
import csv, os, re, sys, collections
sys.stdout.reconfigure(encoding="utf-8"); ROOT = r"J:\VLM-Survey"; os.chdir(ROOT); sys.path.insert(0, ROOT)
_o = sys.stdout; sys.stdout = open(os.devnull, "w", encoding="utf-8")
import main_taxonomy as mt
sys.stdout = _o
Q = r"(?<!vector )(?<!vector-)(?<!residual )(?<!product )quantiz(?![^.]{0,40}?(?:tokeniz|codebook|\btokens?\b|\baction))"
KD = r"(?:small|smaller|compact|lightweight|light-weight|tiny|compress|fewer parameters|\d+(?:\.\d+)?\s?[MB]\b)"   # 'efficient' 는 증류 자체를 꾸미는 경우가 많아 제외
DIST = rf"(?:knowledge distillation|teacher[- ]student)[^.]{{0,150}}{KD}|{KD}[^.]{{0,150}}(?:knowledge distillation|teacher[- ]student)"
SP = r"(?:weight|parameter|model|activation)[- ]sparsi|sparsif\w* (?:the )?(?:weights|model|parameters|LLM|network)|(?:N:M|structured|unstructured|semi-structured) sparsity"
NEW43 = (Q + r"|\bPTQ\b|\bQAT\b|\bINT[48]\b|W4A16|W8A8|W4A4|binariz|ternary|" + DIST +
         r"|low[- ]rank (?:compress|decompos|approximat|factoriz)|\bSVD\b|weight pruning|structured pruning|model pruning|" + SP)
DEV = r"(?:Jetson|Raspberry|Snapdragon|smartphones?|mobile phones?|iPhone|Orin|Google Pixel|Pixel \d|NPU|FPGA)\b"   # 기준: 이름 있는 기기 — 'edge device' 는 기기 이름이 아님
NEW45 = (rf"(?:on|using|with) (?:a |an |the |single )?(?:NVIDIA )?{DEV}|(?:deployed|run|runs|running) on (?:a |an |the )?(?:NVIDIA )?{DEV}")
if __name__ == "__main__":
    from openpyxl import load_workbook
    K = list(csv.DictReader(open("main_survey/classval/answer_key_precision.csv", encoding="utf-8-sig")))
    A = {r[0]: r for r in load_workbook("main_survey/classval/returned_평가자1.xlsx", data_only=True)["A_정밀도"].iter_rows(min_row=2, values_only=True) if r[0]}
    by = {p["id"]: p for p in mt.papers}
    old = {b: re.compile(l, re.I) for b, _, l in mt.BRANCHES}
    n43 = re.compile(NEW43, re.I); s43 = re.compile(mt.PROP + "(?:" + NEW43 + ")", re.I); n45s = re.compile(NEW45, re.I); o45s = re.compile(mt.STRICT_OVERRIDE["4.5"], re.I)
    o43s = re.compile(mt.PROP + "(?:" + mt.BRANCHES[2][2] + ")", re.I)
    def hits(b, p, kind, new):
        if b == "4.3":
            L = n43 if new else old["4.3"]; S = s43 if new else o43s
            return bool(L.search(p["t"])) if kind == "loose" else bool(S.search(p["t"]) or L.search(p["title"]))
        L = old["4.5"]; S = n45s if new else o45s
        return bool(L.search(p["t"])) if kind == "loose" else bool(S.search(p["t"]))
    for b in ("4.3", "4.5"):
        for y in ("2024", "2026"):
            for kind, qi in (("strict", 8), ("loose", 7)):
                for new in (False, True):
                    R = [A[int(k["row"])] for k in K if k["branch"] == b and k["year"] == y and hits(b, by[k["arxiv_id"]], kind, new)]
                    print(b, y, kind, "new" if new else "old", f"labelled matches {len(R)}, yes {sum(r[qi] == '예' for r in R)}")
        for kind in ("loose", "strict"):
            for new in (False, True):
                c = collections.Counter(mt.yr(p) for p in mt.core if hits(b, p, kind, new)); n = collections.Counter(mt.yr(p) for p in mt.core)
                print(f"  corpus {b} {kind} {'new' if new else 'old'}: " + " ".join(f"{y} {c[y]} ({100 * c[y] / n[y]:.2f}%)" for y in ("2024", "2025", "2026")) + f"  ×{(c['2026'] / n['2026']) / (c['2024'] / n['2024']):.2f}")
