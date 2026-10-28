# -*- coding: utf-8 -*-
"""분류 규칙 v2.1(4.3·4.5 수정) 재검증용 새 표본(2026-10-14, 안 1). 규칙을 고칠 때 본 표본(1차 300행)과 겹치지 않는 논문만 뽑는다.
   4.3: 연도(2024, 2026) × {엄격 13, 느슨만 12}.  4.5: 엄격만 바뀌었으므로 엄격 층만 — 2024 는 남은 전부(코퍼스에 4편뿐), 2026 은 13편.
   질문·기준·형식은 1차 시트 A 와 같다. 정답(층)은 answer_key_round2.csv 에만. 시드 20261016.
   출력: 분류검증_새표본_평가자1.xlsx, answer_key_round2.csv"""
import csv, os, random, sys
sys.stdout.reconfigure(encoding="utf-8")
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.abspath(os.path.join(HERE, "..", "..")); sys.path.insert(0, ROOT); os.chdir(ROOT)
_o = sys.stdout; sys.stdout = open(os.devnull, "w", encoding="utf-8")
import main_taxonomy as mt
sys.stdout = _o
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
seen = {(k["branch"], k["arxiv_id"]) for k in csv.DictReader(open(os.path.join(HERE, "answer_key_precision.csv"), encoding="utf-8-sig"))}
src = load_workbook(os.path.join(HERE, "분류검증_시트_평가자1.xlsx"))
CRIT = {r[0][:3]: r for r in src["기준"].iter_rows(min_row=2, values_only=True) if r[0]}
rng = random.Random(20261016); rows, key = [], []
RX = {b: (loose, strict) for b, _, loose, strict, _ in mt.RX}
for b in ("4.3", "4.5"):
    loose, strict = RX[b]
    S = {p["id"] for p in mt.core if (strict.search(p["t"]) if b in mt.STRICT_OVERRIDE else (strict.search(p["t"]) or loose.search(p["title"])))}
    L = {p["id"] for p in mt.core if loose.search(p["t"])}
    for y in ("2024", "2026"):
        sp = sorted(p["id"] for p in mt.core if mt.yr(p) == y and p["id"] in S and (b, p["id"]) not in seen)
        lp = sorted(p["id"] for p in mt.core if mt.yr(p) == y and p["id"] in L and p["id"] not in S and (b, p["id"]) not in seen)
        plan = [("strict", sp, 13)] + ([("loose_only", lp, 12)] if b == "4.3" else [])
        for st, pool, k in plan:
            nS = sum(1 for p in mt.core if mt.yr(p) == y and p["id"] in (S if st == "strict" else L - S))
            for x in rng.sample(pool, min(k, len(pool))):
                rows.append((b, x)); key.append(dict(branch=b, arxiv_id=x, year=y, stratum=st, n_stratum=nS, n_unlabelled=len(pool), loose_hit=x in L, strict_hit=x in S))
order = list(range(len(rows))); rng.shuffle(order); order.sort(key=lambda i: rows[i][0]); rows = [rows[i] for i in order]; key = [key[i] for i in order]
for i, k in enumerate(key, 1): k["row"] = i
with open(os.path.join(HERE, "answer_key_round2.csv"), "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=list(key[0])); w.writeheader(); w.writerows(key)
by = {p["id"]: p for p in mt.papers}
thin = Side(style="thin", color="BFBFBF")
wb = Workbook(); g = wb.active; g.title = "안내"
for line in [f"분류 규칙 재검증(새 표본 {len(rows)}편) — 1차 시트 A 와 같은 방법입니다. 이번엔 4.3 압축, 4.5 배포만 있습니다.", "",
             "각 행에 적힌 분야 하나에 대해서만, 제목과 초록을 보고 판단합니다.",
             "  Q1 다루는가 — 이 분야를 의도한 뜻으로 다루는가? (D열)  →  예/아니오",
             "  Q2 엄격 — 더 엄격한 기준(E열)을 만족하는가?  →  예/아니오",
             "지난번에 본 논문은 들어 있지 않습니다. 지난번 판정이나 규칙 결과는 보지 않고 기준대로만 판단해 주세요. answer_key 파일은 열지 않습니다.",
             "예상 시간: 40분 안팎. 다 하면 파일 경로를 알려 주세요."]:
    g.append([line])
g.column_dimensions["A"].width = 130
k = wb.create_sheet("기준"); k.append(["분야", "Q1 (다루는가)", "Q2 (엄격)"])
for b in ("4.3", "4.5"): k.append(list(CRIT[b]))
a = wb.create_sheet("A_정밀도"); a.append(["번호", "분야", "arXiv", "Q1 기준", "Q2 기준", "제목", "초록", "Q1 다루는가", "Q2 엄격", "메모"])
for i, (b, x) in enumerate(rows, 1):
    p = by[x]; c = CRIT[b]; a.append([i, c[0], x, c[1], c[2], " ".join(p["title"].split()), " ".join(p["abstract"].split()), "", "", ""])
for ws, widths, h in ((k, [26, 80, 80], 60), (a, [6, 14, 12, 30, 30, 30, 90, 10, 10, 20], 200)):
    for j, wd in enumerate(widths, 1): ws.column_dimensions[chr(64 + j)].width = wd
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E7E6E6")
    for r in range(2, ws.max_row + 1): ws.row_dimensions[r].height = h
for r in range(2, a.max_row + 1):
    for c in (8, 9): a.cell(r, c).fill = PatternFill("solid", fgColor="FFF2CC")
dv = DataValidation(type="list", formula1='"예,아니오"', allow_blank=True); a.add_data_validation(dv); dv.add(f"H2:I{a.max_row}"); a.freeze_panes = "D2"
wb.active = 0; wb.save(os.path.join(HERE, "분류검증_새표본_평가자1.xlsx"))
import collections
print(len(rows), "rows"); [print(" ", kk, v) for kk, v in sorted(collections.Counter((k["branch"], k["year"], k["stratum"]) for k in key).items())]
