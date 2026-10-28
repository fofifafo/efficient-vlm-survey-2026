# -*- coding: utf-8 -*-
"""외부 리뷰(10/25) 대응 — 실행(4.4) 규칙 민감도. 새 판정 없음.
   문제: 엄격 실행 규칙 재현율이 2024 5/10, 2026 8/11(1차 평가자 라벨)로 해마다 다를 수 있고, 규칙이 '명사→동사' 순서만 잡아
         PyramidDrop·VTW·VoCo-LLaMA 같은 대표 논문을 놓친다(dropping image tokens, compress vision tokens, withdraw …).
   변형 BROAD: 기존 4.4 어휘 + 동사→명사 순서, drop/discard/remove/withdraw/condense/sparsify, 'fewer visual tokens', 'token budget'.
   출력: (1) 두 코퍼스에서 BROAD 의 느슨·엄격 비율, 상대 성장, Fisher(Bonferroni 0.05/12) — 기준과 나란히,
         (2) 재현율 표본(효율 풀 100편, 1차 평가자)에서 기준·BROAD 의 연도별 재현율, (3) 대표 논문 6편이 잡히는지,
         (4) 참고: VLM 판별에 'large multimodal model(s)/LMM' 을 더하면 늘어나는 논문 수.
   출력 파일: main_survey/exec_rule_sensitivity.json. 실행: python exec_rule_sensitivity.py"""
import collections, csv, json, os, re, sys
sys.stdout.reconfigure(encoding="utf-8"); os.chdir(os.path.dirname(os.path.abspath(__file__))); sys.path.insert(0, ".")
_o = sys.stdout; sys.stdout = open(os.devnull, "w", encoding="utf-8")
import main_taxonomy as mt
sys.stdout = _o
from strict_tags import fisher_exact
from openpyxl import load_workbook
BASE44 = next(l for b, _, l in mt.BRANCHES if b == "4.4")
V = r"(?:visual|vision|image|video|patch)\s+tokens?"
ACT = r"(?:prun|reduc|compress|merg|drop|discard|remov|withdr[ae]w|condens|sparsif|select|evict)"
BROAD44 = (BASE44 + rf"|{ACT}\w*\s+(?:\w+\s+){{0,4}}?{V}|{V}\s+(?:\w+\s+){{0,3}}?{ACT}\w*|fewer\s+{V}|(?:visual|vision|image)\s+token\s+budget"
           r"|token\s+(?:dropping|drop|eviction|sparsification|condensation)")
RULES = {"base": BASE44, "broad": BROAD44}
def matchers(loose):
    lt = re.compile(loose, re.I); st = re.compile(mt.PROP + "(?:" + loose + ")", re.I)
    return (lambda p: bool(lt.search(p["t"]))), (lambda p: bool(st.search(p["t"]) or lt.search(p["title"])))
out = {"rules": {"base": BASE44, "broad": BROAD44}, "trend": [], "recall": {}, "named": {}}
for P, cname in ((mt.pool, "pool"), (mt.core, "vlm")):
    N = collections.Counter(mt.yr(p) for p in P)
    for rk, loose in RULES.items():
        fl, fs = matchers(loose)
        for kind, f in (("loose", fl), ("strict", fs)):
            c = collections.Counter(mt.yr(p) for p in P if f(p)); a, z = c["2024"], c["2026"]
            s24, s26 = 100 * a / N["2024"], 100 * z / N["2026"]; pf = fisher_exact([[a, N["2024"] - a], [z, N["2026"] - z]])[1]
            out["trend"].append(dict(corpus=cname, rule=rk, kind=kind, n24=a, n25=c["2025"], n26=z, s24=round(s24, 2), s26=round(s26, 2),
                                     rel_growth=round(s26 / s24, 2), p_fisher=float(f"{pf:.2e}"), sig_bonf=bool(pf < 0.05 / 12)))
# 재현율(1차 평가자 라벨, 기존 판정만)
codes = lambda v: set() if "없" in str(v or "") else set(re.findall(r"4\.[1-6]", str(v or "")))
KR = list(csv.DictReader(open("main_survey/classval/answer_key_recall.csv", encoding="utf-8-sig")))
G = {int(r[0]): (codes(r[4]), codes(r[5])) for r in load_workbook("main_survey/classval/returned_평가자1.xlsx", data_only=True)["B_재현율"].iter_rows(min_row=2, values_only=True) if r[0]}
PP = {p.get("id") or p.get("arxiv_id"): p for p in mt.papers}
for rk, loose in RULES.items():
    fl, fs = matchers(loose)
    for lv, kind, f in ((0, "loose", fl), (1, "strict", fs)):
        for y in ("2024", "2026", "all"):
            pos = [k for k in KR if (y == "all" or k["year"] == y) and "4.4" in G[int(k["row"])][lv]]
            hit = [k for k in pos if f(PP[k["arxiv_id"]])]
            out["recall"][f"{rk}.{kind}.{y}"] = [len(hit), len(pos)]
# 대표 논문
NAMED = {"2410.17247": "PyramidDrop", "2405.05803": "VTW", "2406.12275": "VoCo-LLaMA", "2407.02392": "TokenPacker", "2412.01818": "FasterVLM/VisPruner", "2403.15388": "LLaVA-PruMerge"}
for aid, nm in NAMED.items():
    p = PP.get(aid)
    if not p: out["named"][nm] = "not in corpus"; continue
    r = {}
    for rk, loose in RULES.items():
        fl, fs = matchers(loose); r[rk] = f"loose={fl(p)} strict={fs(p)}"
    r["in_pool"] = p in mt.pool; out["named"][nm] = r
# 참고: VLM 판별에 LMM 용어 추가
LMM = re.compile(r"large multimodal models?|\bLMMs?\b|large multi-modal models?", re.I)
extra = [p for p in mt.papers if p not in mt.core and LMM.search(p["t"])]
out["lmm_extra_papers"] = dict(n=len(extra), by_year=dict(collections.Counter(mt.yr(p) for p in extra)))
json.dump(out, open("main_survey/exec_rule_sensitivity.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
for t in out["trend"]: print(f"{t['corpus']:4} {t['rule']:5} {t['kind']:6} {t['n24']:3}/{t['n25']:3}/{t['n26']:3}  {t['s24']:5}→{t['s26']:5}  ×{t['rel_growth']}  p={t['p_fisher']}  {'*' if t['sig_bonf'] else ''}")
for k, v in out["recall"].items(): print("recall", k, f"{v[0]}/{v[1]}")
for k, v in out["named"].items(): print("named", k, v)
print("LMM extra:", out["lmm_extra_papers"])
